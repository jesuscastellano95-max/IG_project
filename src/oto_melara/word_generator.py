# -*- coding: utf-8 -*-
"""
Created on Tue Apr 14 21:23:49 2026

@author: jesus
"""
# Librería para comprobar si existe la plantilla Word
import os

# Librería para abrir y editar documentos Word
from docx import Document

# Importamos rutas del proyecto
from common.paths import OTO_WORD_TEMPLATE_PATH, OUTPUT_DIR

# Importamos funciones auxiliares reutilizables
from common.utils import set_cell_text, formatear_numero


# =====================================================
# GENERAR WORD DE OTO MELARA A PARTIR DE RESULTADOS
# =====================================================
def generar_word_oto(data: dict) -> str:
    """
    Genera el documento Word de OTO MELARA
    a partir del diccionario de resultados calculados.

    Devuelve la ruta completa del archivo generado.
    """

    # =====================================================
    # COMPROBAR QUE EXISTE LA PLANTILLA WORD
    # =====================================================
    if not os.path.exists(OTO_WORD_TEMPLATE_PATH):
        raise FileNotFoundError(f"No existe la plantilla Word:\n{OTO_WORD_TEMPLATE_PATH}")

    # =====================================================
    # ABRIR PLANTILLA Y SELECCIONAR LA PRIMERA TABLA
    # =====================================================
    doc = Document(OTO_WORD_TEMPLATE_PATH)
    table = doc.tables[0]

    # =====================================================
    # FILA DE VELOCIDAD
    # row = 2
    # =====================================================
    set_cell_text(
        table.cell(2, 1),
        f"V0c ∈ [{formatear_numero(data['r1_min'])}, {formatear_numero(data['r1_max'])}] m/s"
    )
    set_cell_text(
        table.cell(2, 2),
        f"V0c ∉ [{formatear_numero(data['r2_min'])}, {formatear_numero(data['r2_max'])}] m/s"
    )
    set_cell_text(
        table.cell(2, 4),
        f"V̄0c = {formatear_numero(data['v_med'])} m/s"
    )
    set_cell_text(table.cell(2, 5), data["res_vel"])

    # =====================================================
    # FILA DE DESVIACIÓN ESTÁNDAR
    # row = 3
    # =====================================================
    set_cell_text(
        table.cell(3, 1),
        f"σV0 ≤ {formatear_numero(data['limite_desv'])} m/s"
    )
    set_cell_text(
        table.cell(3, 2),
        f"σV0 > {formatear_numero(data['limite_desv'])} m/s"
    )
    set_cell_text(
        table.cell(3, 4),
        f"σV0 = {formatear_numero(data['desv'])} m/s"
    )
    set_cell_text(table.cell(3, 5), data["res_desv"])

    # =====================================================
    # FILA DE PRESIÓN MÁXIMA
    # row = 4
    # =====================================================
    set_cell_text(
        table.cell(4, 1),
        f"Pmáx ≤ {formatear_numero(data['pmax_lim'])} bar"
    )
    set_cell_text(
        table.cell(4, 2),
        f"Pmáx > {formatear_numero(data['pmax_lim'])} bar"
    )
    set_cell_text(
        table.cell(4, 4),
        f"Pmax = {formatear_numero(data['pmax'])} bar"
    )
    set_cell_text(table.cell(4, 5), data["res_pmax"])

    # =====================================================
    # FILA DE PRESIÓN MEDIA
    # row = 5
    # =====================================================
    set_cell_text(
        table.cell(5, 1),
        f"Pmed < {formatear_numero(data['pmed_lim'])} bar"
    )
    set_cell_text(
        table.cell(5, 2),
        f"Pmed ≥ {formatear_numero(data['pmed_lim'])} bar"
    )
    set_cell_text(
        table.cell(5, 4),
        f"Pmed = {formatear_numero(data['pmed'])} bar"
    )
    set_cell_text(table.cell(5, 5), data["res_pmed"])

    # =====================================================
    # FILA DE ESPOLETA
    # row = 6
    # =====================================================
    set_cell_text(
        table.cell(6, 4),
        f"{data['fallos_espoleta']} fallos"
    )
    set_cell_text(table.cell(6, 5), data["res_espoleta"])

    # =====================================================
    # FILA DE ESTOPÍN
    # row = 7
    # =====================================================
    set_cell_text(
        table.cell(7, 4),
        f"{data['fallos_estopin']} fallos"
    )
    set_cell_text(table.cell(7, 5), data["res_estopin"])

    # =====================================================
    # GUARDAR DOCUMENTO EN OUTPUT
    # =====================================================
    out_path = os.path.join(OUTPUT_DIR, f"{data['nombre_prueba']}.docx")
    doc.save(out_path)

    # Devolver la ruta del archivo generado
    return out_path
